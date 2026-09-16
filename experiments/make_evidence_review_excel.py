import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation


INPUT = "evidence_benchmark.csv"
OUTPUT = "evidence_benchmark_review.xlsx"


print("Loading benchmark...")

df = pd.read_csv(
    INPUT
)

# ------------------------------------------------------------
# Create Excel
# ------------------------------------------------------------

with pd.ExcelWriter(
    OUTPUT,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Evidence Review",
        index=False
    )

    guide = pd.DataFrame({

        "Field": [
            "relevant",
            "actionable",
            "safe_to_ground"
        ],

        "YES": [

            "Historical customer has the same or very closely related underlying issue.",

            "Historical Verizon response provides a useful troubleshooting step, clarification, or concrete support action.",

            "The historical response can safely inform a new reply without assuming current account/state information."
        ],

        "NO": [

            "Different issue, unrelated service, or lexical similarity without substantive semantic match.",

            "Generic acknowledgement, irrelevant response, or response with no useful next step.",

            "Response contains unsupported historical/current-state assumptions or otherwise should not be reused as grounding."
        ]

    })

    guide.to_excel(
        writer,
        sheet_name="Annotation Guide",
        index=False
    )


# ------------------------------------------------------------
# Formatting
# ------------------------------------------------------------

wb = load_workbook(
    OUTPUT
)

ws = wb[
    "Evidence Review"
]

guide = wb[
    "Annotation Guide"
]

# Header formatting
for cell in ws[1]:

    cell.font = Font(
        bold=True
    )

    cell.alignment = Alignment(
        vertical="top",
        wrap_text=True
    )

# Freeze header
ws.freeze_panes = "A2"

# Enable filters
ws.auto_filter.ref = ws.dimensions

# Column widths
widths = {
    "A": 14,
    "B": 14,
    "C": 45,
    "D": 10,
    "E": 12,
    "F": 45,
    "G": 45,
    "H": 18,
    "I": 14,
    "J": 14,
    "K": 18,
    "L": 40
}

for col, width in widths.items():

    ws.column_dimensions[
        col
    ].width = width

# Wrap text
for row in ws.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

# ------------------------------------------------------------
# Dropdowns
# ------------------------------------------------------------

yes_no = DataValidation(
    type="list",
    formula1='"YES,NO"',
    allow_blank=True
)

ws.add_data_validation(
    yes_no
)

# Columns I, J, K
yes_no.add(
    f"I2:I{ws.max_row}"
)

yes_no.add(
    f"J2:J{ws.max_row}"
)

yes_no.add(
    f"K2:K{ws.max_row}"
)

# Guide formatting
for cell in guide[1]:

    cell.font = Font(
        bold=True
    )

    cell.alignment = Alignment(
        vertical="top",
        wrap_text=True
    )

for col in ["A", "B", "C"]:

    guide.column_dimensions[
        col
    ].width = 35

for row in guide.iter_rows():

    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

guide.freeze_panes = "A2"

wb.save(
    OUTPUT
)

print(
    "Created:",
    OUTPUT
)