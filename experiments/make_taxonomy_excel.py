import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

INPUT = "taxonomy_pilot.csv"
OUTPUT = "taxonomy_pilot.xlsx"

df = pd.read_csv(INPUT)

with pd.ExcelWriter(OUTPUT, engine="openpyxl") as writer:
    df.to_excel(writer, index=False, sheet_name="Taxonomy Pilot")

wb = load_workbook(OUTPUT)
ws = wb["Taxonomy Pilot"]

# Freeze header
ws.freeze_panes = "A2"

# Filter
ws.auto_filter.ref = ws.dimensions

# Header
header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center")

# Column widths
widths = {
    "A": 10,
    "B": 10,
    "C": 14,
    "D": 30,
    "E": 80,
    "F": 25,
    "G": 20,
    "H": 45
}

for col, width in widths.items():
    ws.column_dimensions[col].width = width

# Wrap text
for row in ws.iter_rows(min_row=2):
    for cell in row:
        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

# Highlight fields you need to fill
yellow = PatternFill("solid", fgColor="FFF2CC")

for row in range(2, ws.max_row + 1):
    ws.cell(row, 6).fill = yellow
    ws.cell(row, 7).fill = yellow
    ws.cell(row, 8).fill = yellow

# Intent dropdown
intents = (
    "I01,I02,I03,I04,I05,I06,"
    "I07,I08,I09,I10,I11,I12"
)

intent_dropdown = DataValidation(
    type="list",
    formula1=f'"{intents}"',
    allow_blank=True
)

ws.add_data_validation(intent_dropdown)
intent_dropdown.add(f"F2:F{ws.max_row}")

# Confidence dropdown
confidence_dropdown = DataValidation(
    type="list",
    formula1='"high,medium,low"',
    allow_blank=True
)

ws.add_data_validation(confidence_dropdown)
confidence_dropdown.add(f"G2:G{ws.max_row}")

# Instructions sheet
instructions = wb.create_sheet("Instructions")

instructions["A1"] = "Taxonomy Pilot — Instructions"
instructions["A1"].font = Font(size=16, bold=True)

instructions["A3"] = (
    "Read each customer message and select the PRIMARY support need."
)

instructions["A4"] = (
    "Use the yellow cells in the Taxonomy Pilot sheet."
)

instructions["A5"] = (
    "For confidence choose high, medium, or low."
)

instructions["A6"] = (
    "Use notes when the example is ambiguous or context is missing."
)

instructions["A8"] = "I01"
instructions["B8"] = "internet_outage"

instructions["A9"] = "I02"
instructions["B9"] = "internet_performance"

instructions["A10"] = "I03"
instructions["B10"] = "router_equipment"

instructions["A11"] = "I04"
instructions["B11"] = "tv_channel_service"

instructions["A12"] = "I05"
instructions["B12"] = "mobile_phone_service"

instructions["A13"] = "I06"
instructions["B13"] = "fios_app_auth"

instructions["A14"] = "I07"
instructions["B14"] = "billing_payment"

instructions["A15"] = "I08"
instructions["B15"] = "installation_scheduling"

instructions["A16"] = "I09"
instructions["B16"] = "voice_voicemail"

instructions["A17"] = "I10"
instructions["B17"] = "account_fraud"

instructions["A18"] = "I11"
instructions["B18"] = "service_complaint"

instructions["A19"] = "I12"
instructions["B19"] = "other_unclear"

instructions.column_dimensions["A"].width = 12
instructions.column_dimensions["B"].width = 35

wb.save(OUTPUT)

print()
print("SUCCESS!")
print(f"Created: {OUTPUT}")
print(f"Examples: {len(df)}")
print()
print("Open taxonomy_pilot.xlsx in Excel.")